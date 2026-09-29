import json, re, unicodedata
from datetime import date, datetime
from pathlib import Path
import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CFG = json.loads((ROOT / 'config/sources.json').read_text(encoding='utf-8'))
OUT = ROOT / 'data/evaluaciones.json'
TMP = ROOT / '.tmp'
TMP.mkdir(exist_ok=True)

MONTHS = {
    'ENERO': 1, 'FEBRERO': 2, 'MARZO': 3, 'ABRIL': 4, 'MAYO': 5,
    'JUNIO': 6, 'JULIO': 7, 'AGOSTO': 8, 'SEPTIEMBRE': 9,
    'OCTUBRE': 10, 'NOVIEMBRE': 11, 'DICIEMBRE': 12,
}
WEEKDAYS = ('LUNES', 'MARTES', 'MIERCOLES', 'JUEVES', 'VIERNES', 'SABADO', 'DOMINGO')


def clean(v):
    return '' if v is None else ' '.join(str(v).replace('\n', ' ').split())


def norm(v):
    t = clean(v).upper().replace('º', '°')
    return ''.join(c for c in unicodedata.normalize('NFD', t) if unicodedata.category(c) != 'Mn')


def applies(v, grade, section):
    t = norm(v)
    if not t:
        return False
    if re.search(rf'(?<!\d){grade}\s*°?\s*{section}(?![A-Z])', t):
        return True
    m = re.search(rf'(?<!\d){grade}\s*°?\s*(?:BASICO\s*)?([^.;:\n]{{0,45}})', t)
    if m:
        letters = re.findall(r'(?<![A-Z])[A-D](?![A-Z])', m.group(1))
        if section in letters:
            return True
    return False


def mentions_grade(v, grade):
    t = norm(v)
    if not t:
        return False
    return bool(re.search(rf'(?<!\d){grade}\s*°?', t))


def looks_like_generic_activity(v, grade):
    """Actividad fechada del calendario que no está dirigida a ningún paralelo del nivel."""
    text = clean(v)
    t = norm(v)
    if not text or mentions_grade(v, grade):
        return False
    if len(re.sub(r'[^A-Z]', '', t)) < 4:
        return False
    if month_from_sheet(text):
        return False
    if day_from_header(v) or bare_day(v):
        return False
    if any(t == day or t.startswith(day + ' ') for day in WEEKDAYS):
        return False
    # Evita títulos/leyendas típicas sin contenido de actividad.
    noise = (
        'CALENDARIO', 'EVALUACIONES', 'EVALUACION', 'ASIGNATURA', 'ACTIVIDAD',
        'CURSO', 'FECHA', 'SEMANA', 'OBSERVACIONES', 'OBSERVACION'
    )
    if t in noise:
        return False
    return True


def month_from_sheet(title):
    t = norm(title)
    for name, mo in MONTHS.items():
        if name in t:
            return mo
    return None


def merged_anchor_value(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return ws.cell(row, col).value


def event_column_span(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return range(rng.min_col, rng.max_col + 1)
    return range(col, col + 1)


def event_anchor(ws, row, col):
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return rng.min_row, rng.min_col
    return row, col


def day_from_header(v):
    t = norm(v)
    if not t or not any(day in t for day in WEEKDAYS):
        return None
    m = re.search(r'(?<!\d)([12]?\d|3[01])(?!\d)', t)
    if not m:
        return None
    d = int(m.group(1))
    return d if 1 <= d <= 31 else None


def bare_day(v):
    if isinstance(v, (int, float)) and int(v) == v and 1 <= int(v) <= 31:
        return int(v)
    t = clean(v)
    if re.fullmatch(r'\d{1,2}', t):
        d = int(t)
        return d if 1 <= d <= 31 else None
    return None


def date_for_event(ws, row, col):
    mo = month_from_sheet(ws.title)
    if not mo:
        return None
    cols = list(event_column_span(ws, row, col))
    for r in range(row - 1, max(0, row - 6), -1):
        for c in cols:
            dn = day_from_header(merged_anchor_value(ws, r, c))
            if dn:
                try:
                    return date(2026, mo, dn)
                except ValueError:
                    pass
    for r in range(row - 1, max(0, row - 4), -1):
        for c in cols:
            dn = bare_day(merged_anchor_value(ws, r, c))
            if dn:
                try:
                    return date(2026, mo, dn)
                except ValueError:
                    pass
    return None


def subject_and_description(v):
    text = clean(v)
    body = re.sub(
        r'^\s*\d+\s*[°º]?\s*(?:[A-D](?:\s*[,\-/YAND]+\s*(?:\d+\s*[°º]?\s*)?[A-D])*)\s*[:.-]?\s*',
        '', text, flags=re.IGNORECASE
    ).strip()
    if not body:
        body = text
    m = re.match(r'^([A-Za-zÁÉÍÓÚÑáéíóúñ /&]+?)\s*:\s*(.+)$', body)
    if m:
        subject = clean(m.group(1)).title()
        description = clean(m.group(2))
    else:
        subject = 'Actividad escolar'
        description = body
    return subject, description


def extract(key, src):
    grade = re.match(r'(\d+)', key).group(1)
    section = key[-1]
    label = src['curso']
    generic_label = f'General {grade}°'
    url = f"https://docs.google.com/spreadsheets/d/{src['sheet_id']}/export?format=xlsx"
    path = TMP / f'{key}.xlsx'
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    path.write_bytes(r.content)
    wb = load_workbook(path, data_only=True, read_only=False)
    rows, seen = [], set()

    for ws in wb.worksheets:
        if not ws.max_row or not ws.max_column:
            continue
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                v = ws.cell(row, col).value
                specific = applies(v, grade, section)
                generic = False
                if not specific:
                    d_candidate = date_for_event(ws, row, col)
                    generic = bool(d_candidate and looks_like_generic_activity(v, grade))
                if not specific and not generic:
                    continue

                d = date_for_event(ws, row, col)
                if not d:
                    if specific:
                        print(f'[sin fecha] {label} {ws.title} R{row}C{col}: {clean(v)[:140]}')
                    continue

                subject, description = subject_and_description(v)
                ar, ac = event_anchor(ws, row, col)
                course_label = label if specific else generic_label
                rec = {
                    'fecha': d.isoformat(),
                    'curso': course_label,
                    'nivel': f'{grade}°',
                    'es_general': not specific,
                    'asignatura': subject,
                    'descripcion': description,
                    'texto_original': clean(v),
                    'tipo': 'Actividad general' if generic else 'Calendario oficial',
                    'hoja': ws.title,
                    'origen': f"{key}|{ws.title}|R{ar}C{ac}",
                }
                fp = (rec['fecha'], rec['curso'], norm(rec['texto_original']))
                if fp not in seen:
                    seen.add(fp)
                    rows.append(rec)

    print(label, 'registros:', len([x for x in rows if not x['es_general']]))
    print(generic_label, 'registros:', len([x for x in rows if x['es_general']]))
    return rows


records = []
for key, src in CFG.items():
    records.extend(extract(key, src))

records.sort(key=lambda x: (x['fecha'], x['curso'], x['asignatura']))
payload = {
    'actualizado_en': datetime.now().astimezone().strftime('%d-%m-%Y %H:%M'),
    'fuente': 'Colegio Andrée English School',
    'cursos': ['3°D', '6°B'],
    'evaluaciones': records,
}
OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
print('Total registros publicados:', len(records))
