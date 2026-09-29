from __future__ import annotations

import json
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "config" / "sources.json"
OUTPUT = ROOT / "data" / "evaluaciones.json"
TMP = ROOT / ".tmp"
TMP.mkdir(exist_ok=True)

MONTHS = {
    "ENERO": 1, "FEBRERO": 2, "MARZO": 3, "ABRIL": 4, "MAYO": 5,
    "JUNIO": 6, "JULIO": 7, "AGOSTO": 8, "SEPTIEMBRE": 9,
    "OCTUBRE": 10, "NOVIEMBRE": 11, "DICIEMBRE": 12,
}


def clean(value) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\n", " ").split())


def norm(value) -> str:
    txt = clean(value).upper()
    txt = "".join(c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn")
    txt = txt.replace("º", "°")
    return txt


def applies_to(value, grade: str, section: str) -> bool:
    """Detecta curso explícito o grupos de secciones que incluyen la sección objetivo."""
    text = norm(value)
    if not text:
        return False

    g = re.escape(str(grade))
    s = re.escape(section.upper())

    direct = re.search(rf"(?<!\d){g}\s*°?\s*{s}(?![A-Z])", text)
    if direct:
        return True

    # Ejemplos: 3° A Y D / 6° A, B, C Y D / 6 BASICO B
    grouped = re.search(rf"(?<!\d){g}\s*°?\s*(?:BASICO\s*)?([^.;:\n]{{0,35}})", text)
    if grouped:
        tail = grouped.group(1)
        letters = re.findall(r"(?<![A-Z])[A-D](?![A-Z])", tail)
        if section.upper() in letters:
            return True

    return False


def parse_date_value(value, default_year: int = 2026):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = norm(value)
    if not text:
        return None

    # dd-mm-yyyy / dd/mm/yyyy
    m = re.search(r"\b(\d{1,2})[/-](\d{1,2})(?:[/-](\d{2,4}))?\b", text)
    if m:
        d, mo = int(m.group(1)), int(m.group(2))
        y = int(m.group(3)) if m.group(3) else default_year
        if y < 100:
            y += 2000
        try:
            return date(y, mo, d)
        except ValueError:
            pass

    # 15 DE OCTUBRE [DE 2026]
    for name, mo in MONTHS.items():
        m = re.search(rf"\b(\d{{1,2}})\s+(?:DE\s+)?{name}(?:\s+(?:DE\s+)?(20\d{{2}}))?\b", text)
        if m:
            y = int(m.group(2)) if m.group(2) else default_year
            try:
                return date(y, mo, int(m.group(1)))
            except ValueError:
                pass
    return None


def nearby_date(ws, row: int, col: int):
    # 1) misma fila
    for c in range(1, ws.max_column + 1):
        d = parse_date_value(ws.cell(row, c).value)
        if d:
            return d

    # 2) hacia arriba en la misma columna y columnas vecinas
    for r in range(row - 1, max(0, row - 10), -1):
        for c in range(max(1, col - 2), min(ws.max_column, col + 2) + 1):
            d = parse_date_value(ws.cell(r, c).value)
            if d:
                return d
    return None


def row_text(ws, row: int):
    vals = [clean(ws.cell(row, c).value) for c in range(1, ws.max_column + 1)]
    return [v for v in vals if v]


def guess_subject(parts, course_text):
    ignore = {norm(course_text)}
    candidates = []
    for p in parts:
        n = norm(p)
        if not n or n in ignore:
            continue
        if parse_date_value(p):
            continue
        if len(p) <= 45:
            candidates.append(p)
    return candidates[0] if candidates else "Evaluación"


def extract_source(key: str, source: dict):
    grade = re.match(r"(\d+)", key).group(1)
    section = key[-1].upper()
    course_label = source["curso"]
    sheet_id = source["sheet_id"]
    url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx"
    out = TMP / f"{key}.xlsx"

    print(f"Descargando {course_label}...")
    r = requests.get(url, timeout=45)
    r.raise_for_status()
    out.write_bytes(r.content)

    wb = load_workbook(out, data_only=True, read_only=True)
    records = []
    seen = set()

    for ws in wb.worksheets:
        for row in range(1, ws.max_row + 1):
            for col in range(1, ws.max_column + 1):
                value = ws.cell(row, col).value
                if not applies_to(value, grade, section):
                    continue

                d = nearby_date(ws, row, col)
                if not d:
                    print(f"  [sin fecha] {ws.title}!R{row}C{col}: {clean(value)[:100]}")
                    continue

                parts = row_text(ws, row)
                course_text = clean(value)
                description_parts = [p for p in parts if norm(p) != norm(course_text) and not parse_date_value(p)]
                description = " · ".join(description_parts).strip(" ·")
                subject = guess_subject(parts, course_text)

                rec = {
                    "fecha": d.isoformat(),
                    "curso": course_label,
                    "asignatura": subject,
                    "descripcion": description or course_text,
                    "tipo": "Calendario oficial",
                    "hoja": ws.title,
                }
                fingerprint = (rec["fecha"], rec["curso"], norm(rec["descripcion"]))
                if fingerprint not in seen:
                    seen.add(fingerprint)
                    records.append(rec)

    return records


def main():
    sources = json.loads(CONFIG.read_text(encoding="utf-8"))
    all_records = []
    for key, source in sources.items():
        all_records.extend(extract_source(key, source))

    all_records.sort(key=lambda x: (x["fecha"], x["curso"], x["asignatura"]))
    payload = {
        "actualizado_en": datetime.now().astimezone().strftime("%d-%m-%Y %H:%M"),
        "fuente": "Colegio Andrée English School",
        "cursos": ["3°D", "6°B"],
        "evaluaciones": all_records,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Total registros publicados: {len(all_records)}")
    print(f"Archivo: {OUTPUT}")


if __name__ == "__main__":
    main()
